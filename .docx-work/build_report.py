from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


REFERENCE = Path(r"C:\Users\Admin\Desktop\Отчёт по практике\Отчёт по практике.docx")
OUTPUT = Path(r"C:\Users\Admin\Desktop\AMK-Data-Visualization\.docx-work\Отчёт по практике.docx")
IMAGE_MAIN = Path(r"C:\Users\Admin\AppData\Local\Temp\codex-clipboard-db82d0ee-9614-4ba8-8516-004a3731b2ef.png")
IMAGE_METROLOGY = Path(r"C:\Users\Admin\AppData\Local\Temp\codex-clipboard-28644279-8a84-43a1-8697-82569e7239a5.png")
IMAGE_SCHEMATIC = Path(r"C:\Users\Admin\AppData\Local\Temp\codex-clipboard-fc3e01c7-0361-482f-9eb7-5413deb05372.png")
PENDING_PAGE_BREAK = False


def set_run_font(run, name="Times New Roman", size=14, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def configure_styles(doc: Document):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(13)
    pf = normal.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.first_line_indent = Cm(1.25)
    pf.line_spacing = 1.3
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.widow_control = True

    for name, size, align, before, after in [
        ("Heading 1", 16, WD_ALIGN_PARAGRAPH.CENTER, 0, 8),
        ("Heading 2", 14, WD_ALIGN_PARAGRAPH.LEFT, 8, 4),
        ("Heading 3", 14, WD_ALIGN_PARAGRAPH.LEFT, 6, 3),
    ]:
        try:
            style = styles[name]
        except KeyError:
            style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.alignment = align
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    caption = styles.add_style("AMK Caption", WD_STYLE_TYPE.PARAGRAPH)
    caption.font.name = "Times New Roman"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    caption.font.size = Pt(12)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = Cm(0)
    caption.paragraph_format.line_spacing = 1.0
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(6)
    caption.paragraph_format.keep_with_next = True

    try:
        source = styles["AMK Source"]
    except KeyError:
        source = styles.add_style("AMK Source", WD_STYLE_TYPE.PARAGRAPH)
    source.font.name = "Times New Roman"
    source._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    source._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    source._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    source.font.size = Pt(12)
    source.paragraph_format.line_spacing = 1.15
    source.paragraph_format.space_after = Pt(3)


def clear_paragraph(paragraph):
    for child in list(paragraph._element):
        if child.tag != qn("w:pPr"):
            paragraph._element.remove(child)


def add_heading(doc, text, level=1, paragraph=None):
    global PENDING_PAGE_BREAK
    p = paragraph if paragraph is not None else doc.add_paragraph()
    clear_paragraph(p)
    p.style = doc.styles[f"Heading {level}"]
    r = p.add_run(text)
    set_run_font(r, size=16 if level == 1 else 14, bold=True)
    if PENDING_PAGE_BREAK:
        p.paragraph_format.page_break_before = True
        PENDING_PAGE_BREAK = False
    return p


def add_body(doc, text, *, no_indent=False, italic=False):
    p = doc.add_paragraph(style="Normal")
    if no_indent:
        p.paragraph_format.first_line_indent = Cm(0)
    r = p.add_run(text)
    set_run_font(r, size=13, italic=italic)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph(style="AMK Caption")
    r = p.add_run(text)
    set_run_font(r, size=12)
    return p


def add_page_break(doc):
    global PENDING_PAGE_BREAK
    PENDING_PAGE_BREAK = True


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    set_run_font(run, size=14)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "2"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def configure_footer(doc: Document):
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    footer = section.footer
    p = footer.paragraphs[0]
    clear_paragraph(p)
    add_page_field(p)
    first = section.first_page_footer
    for p in first.paragraphs:
        clear_paragraph(p)

    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_cm):
    cell.width = Cm(width_cm)


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "808080")
        borders.append(node)


def add_table(doc, caption, headers, rows, widths):
    add_caption(doc, caption)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_table_borders(table)
    for col, (header, width) in enumerate(zip(headers, widths)):
        cell = table.rows[0].cells[col]
        set_cell_width(cell, width)
        set_cell_shading(cell, "D9EAF7")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(header)
        set_run_font(r, size=12, bold=True)
    set_repeat_table_header(table.rows[0])
    for row_data in rows:
        row = table.add_row()
        for col, (value, width) in enumerate(zip(row_data, widths)):
            cell = row.cells[col]
            set_cell_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if col else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(value)
            set_run_font(r, size=12)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_picture(doc, image_path, width_cm, alt_text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(0)
    shape = p.add_run().add_picture(str(image_path), width=Cm(width_cm))
    shape._inline.docPr.set("descr", alt_text)
    return shape


def add_toc_line(doc, title, page, level=0):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.75 * level)
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(16), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    r = p.add_run(f"{title}\t{page}")
    set_run_font(r, size=13, bold=(level == 0 and title[:1].isdigit()))


def build():
    shutil.copy2(REFERENCE, OUTPUT)
    doc = Document(OUTPUT)
    configure_styles(doc)
    configure_footer(doc)
    doc.core_properties.title = "Отчёт по практике: разработка ПО визуализации данных приборов Горизонт-ИК"
    doc.core_properties.subject = "Разработка программного обеспечения для визуализации и анализа геофизических данных"
    doc.core_properties.keywords = "Горизонт-ИК, визуализация данных, WinUI 3, метрология, COM"

    # Page 2: contents. The existing first paragraph contains the preserved page break;
    # the second paragraph is the first editable slot after the blank title page.
    add_heading(doc, "СОДЕРЖАНИЕ", level=1, paragraph=doc.paragraphs[1])
    for title, page, level in [
        ("ВВЕДЕНИЕ", 3, 0),
        ("1 АНАЛИЗ ПРЕДМЕТНОЙ ОБЛАСТИ", 4, 0),
        ("1.1 Назначение приборов «Горизонт-ИК»", 4, 1),
        ("1.2 Требования к программному обеспечению", 4, 1),
        ("1.3 Выбор технологической платформы", 4, 1),
        ("2 ПРОЕКТИРОВАНИЕ ПРИЛОЖЕНИЯ", 5, 0),
        ("2.1 Архитектура программного комплекса", 5, 1),
        ("2.2 Представление измерительных данных", 6, 1),
        ("2.3 Работа с файлами измерений", 6, 1),
        ("3 РЕАЛИЗАЦИЯ ВИЗУАЛИЗАЦИИ ДАННЫХ", 7, 0),
        ("3.1 Главное окно и построение графиков", 7, 1),
        ("3.2 Настройка состава отображаемых кривых", 8, 1),
        ("3.3 Масштабирование и навигация по данным", 9, 1),
        ("4 ВЗАИМОДЕЙСТВИЕ С ПРИБОРОМ И МЕТРОЛОГИЯ", 10, 0),
        ("4.1 Обмен данными по последовательному порту", 10, 1),
        ("4.2 Запись результатов измерений", 10, 1),
        ("4.3 Работа с метрологическими параметрами", 11, 1),
        ("5 РЕЗУЛЬТАТЫ РАЗРАБОТКИ", 12, 0),
        ("5.1 Реализованные возможности", 12, 1),
        ("5.2 Проверка работоспособности", 12, 1),
        ("ЗАКЛЮЧЕНИЕ", 13, 0),
        ("СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ", 13, 0),
        ("ПРИЛОЖЕНИЕ А", 14, 0),
    ]:
        add_toc_line(doc, title, page, level)

    # Page 3: introduction.
    add_page_break(doc)
    add_heading(doc, "ВВЕДЕНИЕ")
    add_body(doc, "Производственная практика проходила в компании «АМК Горизонт», выполняющей разработку и сопровождение геофизической аппаратуры и программных средств для обработки результатов измерений. Темой практической работы стала «Разработка ПО визуализации данных приборов “Горизонт-ИК”». Необходимость такого программного обеспечения связана с тем, что измерительная аппаратура формирует большие последовательности многоканальных данных, которые требуется оперативно принимать, сохранять, отображать и анализировать.")
    add_body(doc, "Цель работы состояла в создании настольного приложения, объединяющего средства чтения файлов измерений, построения графиков, настройки набора отображаемых каналов, взаимодействия с прибором по последовательному интерфейсу и работы с метрологическими параметрами. При разработке особое внимание уделялось корректному представлению двоичных данных, удобству сравнения нескольких параметров и возможности наблюдать поступающие кадры в реальном времени.")
    add_body(doc, "Для достижения поставленной цели были определены следующие задачи: разработать модульную структуру программы; описать формат измерительного кадра; реализовать загрузку файлов нескольких типов; подготовить расчётные серии для графиков; создать интерфейс выбора частоты, передатчиков и типа измеряемой величины; обеспечить подключение к COM-порту, контроль целостности пакетов и запись принятых кадров; реализовать чтение и сохранение метрологических данных.")
    add_body(doc, "Объектом исследования являются процессы получения и анализа данных геофизического прибора. Предметом исследования выступают программные методы представления измерительных кадров, построения многосерийных графиков и организации надёжного обмена между прибором и компьютером. Практическая значимость работы заключается в сокращении времени первичного контроля измерений и в предоставлении специалисту единой рабочей среды для просмотра, настройки и сохранения данных.")
    add_body(doc, "В отчёте приведены назначение приложения, обоснование выбранной архитектуры, описание основных алгоритмов и пользовательских сценариев, а также результаты проверки реализованного функционала. В качестве иллюстраций использованы снимки разработанного интерфейса.")

    # Page 4: domain and requirements.
    add_page_break(doc)
    add_heading(doc, "1 АНАЛИЗ ПРЕДМЕТНОЙ ОБЛАСТИ")
    add_heading(doc, "1.1 Назначение приборов «Горизонт-ИК»", level=2)
    add_body(doc, "Приборы семейства «Горизонт-ИК» предназначены для регистрации геофизических сигналов по нескольким передающим и приёмным каналам. Один измерительный кадр содержит идентификационную сигнатуру, номер кадра, температуру и набор значений, связанных с амплитудами, разностями фаз и расчётными параметрами. Измерения выполняются для двух частотных диапазонов, а количество доступных передатчиков определяется конфигурацией конкретного прибора.")
    add_heading(doc, "1.2 Требования к программному обеспечению", level=2)
    add_body(doc, "Приложение должно сохранять точность исходных данных и одновременно представлять их в форме, удобной для инженера. Пользователь должен иметь возможность открыть ранее записанный файл, выбрать необходимые каналы, построить один или два графика, задать диапазон кадров и сохранить интересующий участок. При подключённом приборе программа должна принимать новые кадры без блокировки интерфейса и отображать состояние соединения.")
    add_table(
        doc,
        "Таблица 1.1 – Основные требования к приложению",
        ["№", "Требование", "Назначение"],
        [
            ("1", "Корректное чтение двоичных кадров", "Исключение искажения измерений"),
            ("2", "Многосерийные графики", "Сравнение каналов и частот"),
            ("3", "Работа с COM-портом", "Получение данных в реальном времени"),
            ("4", "Метрологические параметры", "Настройка и расчёт характеристик"),
        ],
        [0.9, 7.1, 8.0],
    )
    add_heading(doc, "1.3 Выбор технологической платформы", level=2)
    add_body(doc, "Для реализации выбраны язык C# и платформа .NET 10. Пользовательский интерфейс создан на WinUI 3, что обеспечивает современные элементы управления Windows, поддержку масштабирования DPI и работу с несколькими окнами. Доступ к последовательному порту выполняется средствами System.IO.Ports, а отрисовка графиков организована на Canvas, что позволяет полностью контролировать координаты, сетку, подписи и линии серий.")

    # Page 5: architecture.
    add_page_break(doc)
    add_heading(doc, "2 ПРОЕКТИРОВАНИЕ ПРИЛОЖЕНИЯ")
    add_heading(doc, "2.1 Архитектура программного комплекса", level=2)
    add_body(doc, "Программный комплекс разделён на функциональные модули. Такое решение уменьшает связанность кода: пользовательский интерфейс не выполняет непосредственное декодирование двоичных структур, файловая подсистема не содержит логики построения графика, а обмен с прибором изолирован от экранных форм. Благодаря этому каждый модуль имеет определённую область ответственности, а изменение одного сценария в меньшей степени влияет на остальные части приложения.")
    add_table(
        doc,
        "Таблица 2.1 – Назначение программных модулей",
        ["Модуль", "Основные функции"],
        [
            ("Core", "Двоичные структуры, сигнатуры, возможности приборов и смысловой доступ к полям"),
            ("IO", "Проверка и загрузка файлов, хранение кадров, экспорт диапазонов"),
            ("Charting", "Конфигурации кривых, подготовка значений и группировка по осям"),
            ("Hardware", "COM-порт, пакетный протокол, контроль CRC, опрос и запись"),
            ("Metrology", "Метрологические данные и расчёт параметров на компьютере"),
            ("App", "Главное окно, дополнительные формы и координация пользовательских сценариев"),
        ],
        [3.2, 12.8],
    )
    add_body(doc, "Центральным объектом прикладного уровня является рабочее пространство графиков. Оно хранит независимые конфигурации первого и второго графиков, состояния масштабирования и ссылки на источники данных. Второй график может использовать тот же набор кадров, что и первый, либо отдельный файл. Такое устройство позволяет сопоставлять различные измерения, не смешивая их в одном хранилище.")
    add_body(doc, "Сервисы создаются при запуске главного окна и передаются дополнительным окнам. Окна настройки подключения, схемы прибора и метрологии являются немодальными: пользователь может переключаться между ними и основным графиком, не теряя введённые значения. При закрытии пользователем окно скрывается, а его состояние сохраняется до завершения приложения.")

    # Page 6: data and files.
    add_page_break(doc)
    add_heading(doc, "2.2 Представление измерительных данных", level=2)
    add_body(doc, "Базовый измерительный кадр представлен структурой фиксированного размера 240 байт. В ней последовательно размещены сигнатура прибора, слово состояния, номер кадра, температура и массивы значений для двух частот и пяти возможных передатчиков. Фиксированная компоновка позволяет считывать кадр без промежуточного текстового преобразования и исключает неоднозначность расположения полей.")
    add_body(doc, "Сигнатура используется для определения типа прибора, числа передатчиков и доступных каналов. После распознавания формируется набор возможностей: наличие пятого передатчика, молчащих каналов, фазовых и амплитудных данных, поддержка файловых контейнеров и опроса по COM-порту. Визуальная часть обращается к измерениям через смысловой слой доступа, поэтому не зависит от физических индексов массивов.")
    add_body(doc, "Для амплитуд предусмотрена нормализация единиц измерения на уровне описания возможностей прибора. Например, значения, передаваемые в формате с фиксированной точкой Q12, преобразуются к единицам, используемым графиком. Такое размещение преобразования в общем слое гарантирует одинаковый результат при чтении файла и при получении кадра от прибора.")
    add_heading(doc, "2.3 Работа с файлами измерений", level=2)
    add_body(doc, "Файловая подсистема поддерживает расширения BIN, BIV и DEV. Для BIN и BIV физический кадр совпадает с канонической структурой, а в DEV перед полезной частью находится служебный заголовок. Перед загрузкой проверяются расширение, размер файла, кратность размеру кадра и поддержка обнаруженной сигнатуры. Только после успешной проверки данные добавляются в репозиторий.")
    add_body(doc, "Репозиторий хранит канонические записи, метаданные каждого кадра, имя исходного файла и номер версии данных. Версия изменяется при загрузке, очистке или добавлении кадра и используется окнами настройки для обновления доступных каналов. Для выбранного диапазона можно сформировать новый BIN-файл либо сохранить копию исходного файла без указанного участка. При исключении диапазона сохраняется физическая форма исходных кадров.")

    # Page 7: main UI and first figure.
    add_page_break(doc)
    add_heading(doc, "3 РЕАЛИЗАЦИЯ ВИЗУАЛИЗАЦИИ ДАННЫХ")
    add_heading(doc, "3.1 Главное окно и построение графиков", level=2)
    add_body(doc, "Главное окно содержит меню, панель управления диапазоном, область графика и строку состояния. Двойной режим делит рабочую область на два независимых графика. Для каждой области задаются частота, состав кривых, диапазон по оси X и собственная история масштабирования. На рисунке 3.1 показан пример одновременного отображения амплитуд и симметризованных фаз.")
    add_picture(doc, IMAGE_MAIN, 16.0, "Главное окно приложения с двумя графиками измерительных данных")
    add_caption(doc, "Рисунок 3.1 – Главное окно приложения в двойном режиме")
    add_body(doc, "Графический компонент получает подготовленные серии, подписи и цвета. Он рассчитывает диапазоны осей, добавляет пятипроцентный отступ от минимального и максимального значений, строит сетку и формирует шаг шкалы по последовательности 1–2–5–10. Для фаз используются миллиградусы, для амплитуд – милливольты, а ось X может отображать номер кадра или температуру.")

    # Page 8: schematic and second figure.
    add_page_break(doc)
    add_heading(doc, "3.2 Настройка состава отображаемых кривых", level=2)
    add_body(doc, "Состав графика задаётся в окне «Схема прибора». В верхней части расположены передатчики T1–T5, приёмники R1 и R2 и элементы молчащего режима. Недоступные для конкретной конфигурации передатчики блокируются автоматически. Выбранные элементы выделяются жёлтой рамкой, после чего программа формирует перечень совместимых кривых.")
    add_picture(doc, IMAGE_SCHEMATIC, 16.0, "Окно выбора передатчиков, приёмников и быстрых пресетов")
    add_caption(doc, "Рисунок 3.2 – Настройка схемы прибора и состава кривых")
    add_body(doc, "Пользователю доступны быстрые конфигурации: амплитуды приёмников, отношение амплитуд, разность фаз, удельное электрическое сопротивление, симметризованные фазы, параметры молчащего режима и температура. Кривые могут быть распределены по левой и правой осям Y. При выборе несовместимого сочетания применение конфигурации блокируется, что предотвращает построение некорректного набора данных.")

    # Page 9: navigation.
    add_page_break(doc)
    add_heading(doc, "3.3 Масштабирование и навигация по данным", level=2)
    add_body(doc, "Для анализа длинных записей реализовано несколько способов навигации. Диапазон по оси X можно задать номерами начального и конечного кадров. Пустое конечное значение означает выбор до последней доступной записи. После применения диапазона горизонтальная полоса прокрутки перемещает окно фиксированной ширины по всему файлу.")
    add_body(doc, "Масштабирование мышью выполняется выделением прямоугольной области графика. Координаты выделения преобразуются в диапазон кадров и значения левой и правой осей. Перед каждым изменением сохраняется снимок состояния, поэтому кнопка отмены возвращает предыдущий масштаб пошагово. Отдельный сброс удаляет пользовательские ограничения и восстанавливает автоматический диапазон.")
    add_body(doc, "Режим «Поток» предназначен для наблюдения за поступающими кадрами. На экране отображаются последние N записей, а окно автоматически перемещается вслед за новыми данными. Горизонтальное масштабирование в этом режиме отключено, поскольку ось X должна следовать за потоком, однако пользователь может изменять диапазон Y и прокручивать его по вертикали.")
    add_body(doc, "Каждая серия выравнивается по индексам кадров. Если значение недоступно для конкретного типа прибора или выходит за допустимый диапазон, в массив помещается специальное значение NaN. При отрисовке оно разрывает линию и тем самым не создаёт ложного соединения между удалёнными точками. Такой подход особенно важен для файлов с неоднородными или временно отсутствующими каналами.")
    add_body(doc, "Легенда формируется с учётом выбранной конфигурации. Для быстрых пресетов название категории выводится один раз, а серии обозначаются короткими именами T1–T5. В комбинированном режиме подписи группируются по осям и типам данных; группы левой и правой осей разделяются визуальным маркером. Перенос легенды на несколько строк исключает наложение надписей при большом числе кривых.")

    # Page 10: serial and recording.
    add_page_break(doc)
    add_heading(doc, "4 ВЗАИМОДЕЙСТВИЕ С ПРИБОРОМ И МЕТРОЛОГИЯ")
    add_heading(doc, "4.1 Обмен данными по последовательному порту", level=2)
    add_body(doc, "Подключение к прибору выполняется через выбранный COM-порт со скоростью 125000 бит/с, восемью битами данных, без контроля чётности и с одним стоповым битом. В окне подключения отображаются состояние соединения, номер последнего кадра, количество принятых пакетов и ошибок, параметры записи и интервал опроса.")
    add_body(doc, "Сервис обмена поддерживает два варианта пакетов. В основном варианте пакет содержит код команды, общую длину, полезную нагрузку и контрольную сумму CRC16. Для отдельного режима прибора используется ответ фиксированной длины: после кода команды находятся два служебных байта и 240 байт измерительного кадра. Выбор варианта выполняется автоматически по сигнатуре и числу передатчиков.")
    add_body(doc, "Контрольная сумма рассчитывается по алгоритму CRC16 с начальным значением 0xFFFF и полиномом 0xA001. После приёма дополнительно проверяются код ответа, длина полезной нагрузки и сигнатура кадра. Ошибки не приводят к остановке приложения: ведётся счётчик последовательных отказов, а после превышения порога порт закрывается и выполняется попытка автоматического переподключения.")
    add_heading(doc, "4.2 Запись результатов измерений", level=2)
    add_body(doc, "Каждый успешно принятый кадр передаётся в основной репозиторий, отображается на связанных графиках и при включённой записи сохраняется в BIN-файл. Пользователь может создать новый файл либо дописывать существующий. По умолчанию используется папка AMK_Data на рабочем столе; выбранный пользователем путь сохраняется в локальных настройках программы.")
    add_body(doc, "Запись выполняется в том же каноническом формате по 240 байт на кадр. Перед дозаписью проверяются существование файла и кратность его размера длине записи. В журнал протокола выводятся переданные и принятые пакеты, размеры фрагментов, результаты CRC, таймауты, выбранный вариант транспорта и причины отклонения данных. Это упрощает диагностику соединения с физическим прибором.")

    # Page 11: metrology and third figure.
    add_page_break(doc)
    add_heading(doc, "4.3 Работа с метрологическими параметрами", level=2)
    add_body(doc, "Окно метрологии предназначено для просмотра и редактирования геометрии зонда, частот, воздушных поправок и идентификационной информации. Метрологическая структура также занимает 240 байт и может быть считана из прибора, загружена из файла, сохранена на диск или записана обратно в зонд с последующей проверкой.")
    add_picture(doc, IMAGE_METROLOGY, 15.4, "Окно метрологии с геометрией зонда и воздушными калибровками")
    add_caption(doc, "Рисунок 4.1 – Окно работы с метрологическими параметрами")
    add_body(doc, "Поля L1, L2, F1 и F2 по умолчанию защищены от случайного редактирования. Сигнатура формируется из типа прибора, числа передатчиков, модификации и номера. На основании метрологических данных программа может рассчитывать симметризованные фазы и удельное электрическое сопротивление на компьютере. Если сигнатуры измерительного кадра и метрологии не совпадают, используются значения, содержащиеся в самом кадре.")

    # Page 12: results and validation.
    add_page_break(doc)
    add_heading(doc, "5 РЕЗУЛЬТАТЫ РАЗРАБОТКИ")
    add_heading(doc, "5.1 Реализованные возможности", level=2)
    add_body(doc, "В результате практики разработано настольное приложение, обеспечивающее полный базовый цикл работы с измерительными данными: открытие файлов, распознавание конфигурации прибора, подготовку и построение серий, выбор каналов, масштабирование, подключение к COM-порту, приём и запись кадров, а также работу с метрологическими параметрами.")
    add_table(
        doc,
        "Таблица 5.1 – Результаты реализации",
        ["Подсистема", "Полученный результат"],
        [
            ("Файлы", "Загрузка BIN, BIV и DEV, экспорт и исключение диапазонов"),
            ("Графики", "Один или два источника, две оси Y, номер кадра или температура по X"),
            ("Навигация", "Диапазон, drag-zoom, отмена, полосы прокрутки и режим «Поток»"),
            ("Прибор", "Поиск COM-портов, опрос, проверка пакетов, reconnect и запись"),
            ("Метрология", "Файл, чтение/запись зонда и расчёт параметров на компьютере"),
        ],
        [4.0, 12.0],
    )
    add_heading(doc, "5.2 Проверка работоспособности", level=2)
    add_body(doc, "Проверка проводилась на уровне сборки решения, запуска окон и выполнения основных пользовательских сценариев. Проект собирается в конфигурации Debug x64 без предупреждений и ошибок. Проверены открытие главного окна, схема прибора, настройки подключения и окно метрологии, переключение двух источников графиков, применение диапазона и операции масштабирования.")
    add_body(doc, "Журнал физического обмена содержит успешные сеансы подключения, автоматический выбор фиксированного варианта протокола и последовательное получение кадров с корректной сигнатурой. Дополнительно контролировались размеры двоичных структур, обработка метаданных и сохранение принятых записей. Для окончательной оценки точности расчётов требуется сравнивать значения с эталонными измерениями на одинаковых исходных данных.")
    add_body(doc, "Основным направлением дальнейшего развития является оптимизация отображения очень больших файлов: прореживание точек и кэширование подготовленных серий. Также целесообразно расширить средства интерактивного анализа координатами курсора и маркерами. Эти задачи не изменяют разработанную архитектуру и могут быть реализованы как отдельные компоненты.")

    # Page 13: conclusion and sources.
    add_page_break(doc)
    add_heading(doc, "ЗАКЛЮЧЕНИЕ")
    add_body(doc, "В ходе производственной практики в компании «АМК Горизонт» было разработано программное обеспечение визуализации данных приборов «Горизонт-ИК». Выполнен анализ предметной области, определены требования к хранению и отображению многоканальных измерений, выбрана технологическая платформа и спроектирована модульная структура приложения.")
    add_body(doc, "Реализованы двоичная модель измерительного кадра, проверка и загрузка файлов, репозиторий данных, подготовка расчётных серий и собственный компонент построения графиков. Пользователь может выбирать частоту и каналы, применять готовые конфигурации, работать с двумя графиками, изменять диапазон, масштабировать данные и наблюдать поступление новых кадров в режиме «Поток».")
    add_body(doc, "Отдельным результатом стала реализация взаимодействия с прибором: поиск COM-портов, обмен пакетами с проверкой CRC16, автоматическое восстановление связи, запись кадров и подробное журналирование. Окно метрологии обеспечивает работу с геометрией зонда, воздушными поправками, идентификацией и расчётными параметрами.")
    add_body(doc, "Поставленная цель достигнута: создано работоспособное приложение, которое объединяет основные операции инженера по просмотру, контролю и сохранению геофизических данных. Полученные навыки включают проектирование архитектуры, работу с двоичными форматами и последовательным портом, разработку интерфейсов WinUI 3 и реализацию алгоритмов визуализации.")
    add_heading(doc, "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ")
    sources = [
        "Microsoft. Документация по Windows App SDK [Электронный ресурс]. URL: https://learn.microsoft.com/windows/apps/windows-app-sdk/ (дата обращения: 16.07.2026).",
        "Microsoft. Обзор WinUI 3 [Электронный ресурс]. URL: https://learn.microsoft.com/windows/apps/winui/winui3/ (дата обращения: 16.07.2026).",
        "Microsoft. Документация по .NET 10 [Электронный ресурс]. URL: https://learn.microsoft.com/dotnet/ (дата обращения: 16.07.2026).",
        "Microsoft. Класс SerialPort [Электронный ресурс]. URL: https://learn.microsoft.com/dotnet/api/system.io.ports.serialport (дата обращения: 16.07.2026).",
        "Технические материалы компании «АМК Горизонт» по форматам данных и протоколам приборов «Горизонт-ИК». Внутренняя документация предприятия.",
    ]
    for source in sources:
        p = doc.add_paragraph(style="AMK Source")
        p.paragraph_format.first_line_indent = Cm(0)
        r = p.add_run(source)
        set_run_font(r, size=12)

    # Page 14: appendix.
    add_page_break(doc)
    add_heading(doc, "ПРИЛОЖЕНИЕ А")
    add_caption(doc, "(обязательное)")
    add_heading(doc, "Структура основных программных компонентов", level=2)
    add_body(doc, "В таблице А.1 приведены ключевые компоненты разработанного приложения и их роль в пользовательских сценариях. Перечень отражает фактическое разделение ответственности между хранением, вычислениями, аппаратным обменом и отображением.")
    add_table(
        doc,
        "Таблица А.1 – Ключевые программные компоненты",
        ["Компонент", "Назначение"],
        [
            ("ToolDataAdapterRegistry", "Определение типа прибора и преобразование кадра к общей модели"),
            ("ToolDataFileReader", "Проверка файла и чтение физических кадров"),
            ("ToolDataRepository", "Хранение записей, метаданных и версии набора данных"),
            ("ChartDataCalculator", "Расчёт амплитудных, фазовых, температурных и метрологических серий"),
            ("ChartWorkspace", "Независимые источники, конфигурации и состояния двух графиков"),
            ("LineChartControl", "Сетка, оси, легенда, линии, выделение и полосы прокрутки"),
            ("SerialDeviceService", "Подключение, опрос, контроль пакетов, reconnect и запись"),
            ("MetrologyCalculator", "Симметризация фаз и вычисление удельного сопротивления"),
        ],
        [6.0, 10.0],
    )
    add_body(doc, "Обобщённый алгоритм обработки данных имеет следующий вид: входной файл или пакет прибора проходит проверку размера и сигнатуры; кадр преобразуется к общей структуре; вместе с ним формируются метаданные возможностей; запись добавляется в репозиторий; выбранная конфигурация преобразует поля кадра в выровненные серии; графический компонент рассчитывает шкалы и отображает результат. При поступлении нового кадра повторяются только этапы добавления, подготовки видимого диапазона и отрисовки.")

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
