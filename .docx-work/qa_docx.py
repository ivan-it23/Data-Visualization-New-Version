from pathlib import Path
from zipfile import ZipFile
import hashlib
import re

from docx import Document


path = Path(r"C:\Users\Admin\Desktop\AMK-Data-Visualization\.docx-work\Отчёт по практике.docx")
doc = Document(path)
with ZipFile(path) as archive:
    xml = archive.read("word/document.xml").decode("utf-8")
    footer = "".join(
        archive.read(name).decode("utf-8")
        for name in archive.namelist()
        if name.startswith("word/footer") and name.endswith(".xml")
    )

text = "\n".join(p.text for p in doc.paragraphs)
first = doc.paragraphs[0]._p.xml
print(f"sha256={hashlib.sha256(path.read_bytes()).hexdigest()}")
print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} images={len(doc.inline_shapes)}")
print(f"chars={len(text)} page_breaks={xml.count('w:type=\"page\"')}")
print(f"first_page_break={'w:type=\"page\"' in first} first_text={doc.paragraphs[0].text!r}")
print(f"tcW={sorted(set(re.findall(r'<w:tcW[^>]*w:w=\"(\d+)\"', xml)), key=int)}")
print(f"footer_page_field={'PAGE' in footer}")
print(f"forbidden_migration={bool(re.search(r'миграц|перенос.{0,20}логик', text, re.I))}")
for i, shape in enumerate(doc.inline_shapes, 1):
    print(f"image_{i}={shape.width / 360000:.2f}x{shape.height / 360000:.2f}cm")
