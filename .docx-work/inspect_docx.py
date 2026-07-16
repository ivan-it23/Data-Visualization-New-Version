from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def inspect(path: Path) -> dict:
    doc = Document(path)
    paragraphs = []
    for index, paragraph in enumerate(doc.paragraphs):
        paragraphs.append(
            {
                "index": index,
                "style": paragraph.style.name if paragraph.style else None,
                "text": paragraph.text,
                "alignment": str(paragraph.alignment),
                "page_break_before": bool(paragraph.paragraph_format.page_break_before),
                "runs": [
                    {
                        "text": run.text,
                        "bold": run.bold,
                        "italic": run.italic,
                        "size_pt": run.font.size.pt if run.font.size else None,
                        "font": run.font.name,
                        "breaks": len(run._element.findall(".//" + qn("w:br"))),
                    }
                    for run in paragraph.runs
                ],
            }
        )

    tables = []
    for table_index, table in enumerate(doc.tables):
        rows = []
        for row in table.rows:
            rows.append([cell.text for cell in row.cells])
        tables.append({"index": table_index, "rows": rows})

    sections = []
    for index, section in enumerate(doc.sections):
        sections.append(
            {
                "index": index,
                "width_mm": round(section.page_width.mm, 2),
                "height_mm": round(section.page_height.mm, 2),
                "left_mm": round(section.left_margin.mm, 2),
                "right_mm": round(section.right_margin.mm, 2),
                "top_mm": round(section.top_margin.mm, 2),
                "bottom_mm": round(section.bottom_margin.mm, 2),
                "header_mm": round(section.header_distance.mm, 2),
                "footer_mm": round(section.footer_distance.mm, 2),
            }
        )

    styles = {}
    for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"]:
        try:
            style = doc.styles[name]
        except KeyError:
            continue
        styles[name] = {
            "font": style.font.name,
            "size_pt": style.font.size.pt if style.font.size else None,
            "bold": style.font.bold,
            "italic": style.font.italic,
        }

    return {
        "path": str(path),
        "paragraphs": paragraphs,
        "tables": tables,
        "sections": sections,
        "styles": styles,
        "inline_shapes": len(doc.inline_shapes),
    }


for argument in sys.argv[1:]:
    print(json.dumps(inspect(Path(argument)), ensure_ascii=False, indent=2))
